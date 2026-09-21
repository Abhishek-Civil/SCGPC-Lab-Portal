import os, io, sqlite3, secrets, zipfile
from datetime import datetime
from functools import wraps
from flask import Flask, render_template, request, redirect, url_for, session, flash, send_file, jsonify
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import load_workbook
from openpyxl.chart import LineChart, BarChart, ScatterChart, Reference
from openpyxl.chart.label import DataLabelList

BASE = os.path.dirname(os.path.abspath(__file__))
DB = os.path.join(BASE, 'lab_portal.db')
TEMPLATE_XLSX = os.path.join(BASE, 'SCGPC_Fire_Resistance_SCBA_template.xlsx')
EXPORT_DIR = os.path.join(BASE, 'exports')

os.makedirs(EXPORT_DIR, exist_ok=True)

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', secrets.token_hex(32))
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024

MIXES = [
    ('M0', 0, '4M', 'Control'),
    ('M1', 5, '4M', 'SCBA 4M'),
    ('M2', 10, '4M', 'SCBA 4M'),
    ('M3', 15, '4M', 'SCBA 4M'),
    ('M4', 20, '4M', 'SCBA 4M'),
    ('M5', 25, '4M', 'SCBA 4M'),
    ('M6', 0, '6M', 'Control'),
    ('M7', 5, '6M', 'SCBA 6M'),
    ('M8', 10, '6M', 'SCBA 6M'),
    ('M9', 15, '6M', 'SCBA 6M'),
    ('M10', 20, '6M', 'SCBA 6M'),
    ('M11', 25, '6M', 'SCBA 6M')
]

TEMPS = [200, 400, 600, 800]
REPS = [1, 2, 3]


def db_conn():
    c = sqlite3.connect(DB)
    c.row_factory = sqlite3.Row
    return c


def init_db():
    c = db_conn()

    c.executescript('''
    CREATE TABLE IF NOT EXISTS users(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        username TEXT UNIQUE NOT NULL,
        password_hash TEXT NOT NULL,
        role TEXT NOT NULL DEFAULT 'member',
        active INTEGER NOT NULL DEFAULT 1,
        created_at TEXT NOT NULL
    );

    CREATE TABLE IF NOT EXISTS mixes(
        mix_id TEXT PRIMARY KEY,
        scba_pct REAL NOT NULL,
        naoh TEXT NOT NULL,
        category TEXT NOT NULL,
        baseline_mpa REAL
    );

    CREATE TABLE IF NOT EXISTS specimens(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        mix_id TEXT NOT NULL,
        temperature INTEGER NOT NULL,
        replicate INTEGER NOT NULL,
        pre_mass REAL,
        pre_condition TEXT,
        post_mass REAL,
        colour TEXT,
        cracks TEXT,
        spalling TEXT,
        failure_load REAL,
        remarks TEXT,
        updated_by TEXT,
        updated_at TEXT,
        UNIQUE(mix_id, temperature, replicate),
        FOREIGN KEY(mix_id) REFERENCES mixes(mix_id)
    );
    ''')

    for m in MIXES:
        c.execute(
            'INSERT OR IGNORE INTO mixes(mix_id,scba_pct,naoh,category) VALUES(?,?,?,?)',
            m
        )

    count = c.execute(
        'SELECT COUNT(*) n FROM specimens'
    ).fetchone()['n']

    if count == 0:
        for mix, *_ in MIXES:
            for t in TEMPS:
                for r in REPS:
                    c.execute(
                        'INSERT INTO specimens(mix_id,temperature,replicate) VALUES(?,?,?)',
                        (mix, t, r)
                    )

    if c.execute('SELECT COUNT(*) n FROM users').fetchone()['n'] == 0:
        u = os.environ.get('ADMIN_USERNAME', 'admin')
        p = os.environ.get('ADMIN_PASSWORD', 'ChangeMe123!')

        c.execute(
            'INSERT INTO users(username,password_hash,role,created_at) VALUES(?,?,?,?)',
            (
                u,
                generate_password_hash(p),
                'admin',
                datetime.utcnow().isoformat()
            )
        )

    c.commit()
    c.close()


# IMPORTANT:
# Initialize the database when Gunicorn starts the application.
init_db()


def login_required(fn):
    @wraps(fn)
    def wrapper(*a, **kw):
        if 'user_id' not in session:
            return redirect(url_for('login'))
        return fn(*a, **kw)

    return wrapper


def admin_required(fn):
    @wraps(fn)
    def wrapper(*a, **kw):
        if session.get('role') != 'admin':
            flash('Admin access required.', 'error')
            return redirect(url_for('dashboard'))

        return fn(*a, **kw)

    return wrapper


def calc_row(row):
    pre = row['pre_mass']
    post = row['post_mass']
    load = row['failure_load']
    base = row['baseline_mpa']

    mass_loss = (
        ((pre - post) / pre * 100)
        if pre is not None and post is not None and pre
        else None
    )

    strength = (
        (load * 1000 / 10000)
        if load is not None
        else None
    )

    residual = (
        (strength / base * 100)
        if strength is not None and base
        else None
    )

    return mass_loss, strength, residual


def fetch_dashboard():
    c = db_conn()

    rows = c.execute('''
        SELECT
            s.*,
            m.scba_pct,
            m.naoh,
            m.category,
            m.baseline_mpa
        FROM specimens s
        JOIN mixes m ON s.mix_id = m.mix_id
        ORDER BY s.temperature, s.mix_id, s.replicate
    ''').fetchall()

    out = []

    for r in rows:
        ml, st, rs = calc_row(r)

        d = dict(r)

        d.update(
            mass_loss=ml,
            strength=st,
            residual=rs
        )

        out.append(d)

    c.close()

    return out


def summary_data():
    rows = fetch_dashboard()
    by = {}

    for r in rows:
        k = (r['mix_id'], r['temperature'])
        by.setdefault(k, []).append(r)

    summary = []

    for mix, scba, naoh, cat in MIXES:
        item = {
            'mix_id': mix,
            'scba': scba,
            'naoh': naoh,
            'category': cat
        }

        for t in TEMPS:
            vals = [
                x['residual']
                for x in by.get((mix, t), [])
                if x['residual'] is not None
            ]

            ml = [
                x['mass_loss']
                for x in by.get((mix, t), [])
                if x['mass_loss'] is not None
            ]

            item[str(t)] = (
                round(sum(vals) / len(vals), 2)
                if vals
                else None
            )

            item[f'ml{t}'] = (
                round(sum(ml) / len(ml), 2)
                if ml
                else None
            )

        summary.append(item)

    return summary


def generate_excel():
    wb = load_workbook(TEMPLATE_XLSX)

    raw = wb['2 - Raw Data Entry']
    mixws = wb['1 - Mix Reference']

    c = db_conn()

    mixes = {
        r['mix_id']: r
        for r in c.execute('SELECT * FROM mixes').fetchall()
    }

    specs = c.execute(
        'SELECT * FROM specimens ORDER BY temperature,mix_id,replicate'
    ).fetchall()

    # Baseline strengths
    for i, mix in enumerate(MIXES, start=4):
        r = mixes[mix[0]]
        mixws.cell(i, 8).value = r['baseline_mpa']

    # Raw data rows 5:148
    for excel_row, s in zip(range(5, 149), specs):
        raw.cell(excel_row, 4).value = s['pre_mass']
        raw.cell(excel_row, 5).value = s['pre_condition']
        raw.cell(excel_row, 6).value = s['post_mass']
        raw.cell(excel_row, 7).value = s['colour']
        raw.cell(excel_row, 8).value = s['cracks']
        raw.cell(excel_row, 9).value = s['spalling']
        raw.cell(excel_row, 10).value = s['failure_load']
        raw.cell(excel_row, 16).value = s['remarks']

    c.close()

    out = os.path.join(
        EXPORT_DIR,
        'SCGPC_Lab_Data_Live.xlsx'
    )

    wb.save(out)

    return out


@app.context_processor
def inject():
    return {
        'current_user': session.get('username'),
        'role': session.get('role')
    }


@app.route('/')
def home():
    return redirect(
        url_for('dashboard')
        if 'user_id' in session
        else url_for('login')
    )


@app.route('/login', methods=['GET', 'POST'])
def login():

    if request.method == 'POST':

        u = request.form.get('username', '').strip()
        p = request.form.get('password', '')

        c = db_conn()

        user = c.execute(
            'SELECT * FROM users WHERE username=? AND active=1',
            (u,)
        ).fetchone()

        c.close()

        if user and check_password_hash(
            user['password_hash'],
            p
        ):
            session.clear()

            session.update(
                user_id=user['id'],
                username=user['username'],
                role=user['role']
            )

            return redirect(url_for('dashboard'))

        flash(
            'Invalid username or password.',
            'error'
        )

    return render_template('login.html')


@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))


@app.route('/dashboard')
@login_required
def dashboard():

    rows = fetch_dashboard()

    complete = sum(
        1
        for r in rows
        if r['failure_load'] is not None
    )

    return render_template(
        'dashboard.html',
        rows=rows,
        summary=summary_data(),
        complete=complete,
        total=len(rows),
        temps=TEMPS
    )


@app.route('/entry')
@login_required
def entry():

    c = db_conn()

    rows = c.execute('''
        SELECT
            s.*,
            m.scba_pct,
            m.naoh,
            m.baseline_mpa
        FROM specimens s
        JOIN mixes m ON s.mix_id = m.mix_id
        ORDER BY s.temperature,s.mix_id,s.replicate
    ''').fetchall()

    c.close()

    return render_template(
        'entry.html',
        rows=rows,
        temps=TEMPS
    )


@app.route('/api/specimen/<int:sid>', methods=['POST'])
@login_required
def save_specimen(sid):

    data = request.get_json(
        silent=True
    ) or request.form

    fields = [
        'pre_mass',
        'pre_condition',
        'post_mass',
        'colour',
        'cracks',
        'spalling',
        'failure_load',
        'remarks'
    ]

    vals = []

    for f in fields:

        v = data.get(f)

        if v == '':
            v = None

        if f in (
            'pre_mass',
            'post_mass',
            'failure_load'
        ) and v is not None:

            try:
                v = float(v)
            except:
                return jsonify(
                    ok=False,
                    error=f'Invalid {f}'
                ), 400

        vals.append(v)

    c = db_conn()

    c.execute('''
        UPDATE specimens
        SET
            pre_mass=?,
            pre_condition=?,
            post_mass=?,
            colour=?,
            cracks=?,
            spalling=?,
            failure_load=?,
            remarks=?,
            updated_by=?,
            updated_at=?
        WHERE id=?
    ''', (
        *vals,
        session['username'],
        datetime.utcnow().isoformat(),
        sid
    ))

    c.commit()

    row = c.execute('''
        SELECT
            s.*,
            m.baseline_mpa
        FROM specimens s
        JOIN mixes m ON s.mix_id=m.mix_id
        WHERE s.id=?
    ''', (sid,)).fetchone()

    c.close()

    ml, st, rs = calc_row(row)

    return jsonify(
        ok=True,
        mass_loss=ml,
        strength=st,
        residual=rs
    )


@app.route('/api/baseline/<mix_id>', methods=['POST'])
@login_required
def save_baseline(mix_id):

    if session.get('role') != 'admin':
        return jsonify(
            ok=False,
            error='Admin only'
        ), 403

    v = request.json.get('baseline_mpa')

    try:
        v = (
            float(v)
            if v not in (None, '')
            else None
        )
    except:
        return jsonify(
            ok=False,
            error='Invalid strength'
        ), 400

    c = db_conn()

    c.execute(
        'UPDATE mixes SET baseline_mpa=? WHERE mix_id=?',
        (v, mix_id)
    )

    c.commit()
    c.close()

    return jsonify(ok=True)


@app.route('/charts')
@login_required
def charts():
    return render_template(
        'charts.html',
        summary=summary_data(),
        temps=TEMPS
    )


@app.route('/results')
@login_required
def results():
    return render_template(
        'results.html',
        summary=summary_data(),
        temps=TEMPS
    )


@app.route('/excel')
@login_required
def excel():

    path = generate_excel()

    return send_file(
        path,
        as_attachment=True,
        download_name='SCGPC_Lab_Data_Live.xlsx'
    )


@app.route('/admin')
@admin_required
def admin():

    c = db_conn()

    users = c.execute(
        'SELECT id,username,role,active,created_at FROM users ORDER BY id'
    ).fetchall()

    mixes = c.execute(
        'SELECT * FROM mixes ORDER BY mix_id'
    ).fetchall()

    c.close()

    return render_template(
        'admin.html',
        users=users,
        mixes=mixes
    )


@app.route('/admin/user', methods=['POST'])
@admin_required
def add_user():

    u = request.form.get(
        'username',
        ''
    ).strip()

    p = request.form.get(
        'password',
        ''
    )

    role = request.form.get(
        'role',
        'member'
    )

    if not u or len(p) < 8:
        flash(
            'Username required and password must be at least 8 characters.',
            'error'
        )

        return redirect(
            url_for('admin')
        )

    c = db_conn()

    try:

        c.execute(
            '''
            INSERT INTO users(
                username,
                password_hash,
                role,
                created_at
            )
            VALUES(?,?,?,?)
            ''',
            (
                u,
                generate_password_hash(p),
                role,
                datetime.utcnow().isoformat()
            )
        )

        c.commit()

        flash(
            'User created.',
            'ok'
        )

    except sqlite3.IntegrityError:

        flash(
            'Username already exists.',
            'error'
        )

    c.close()

    return redirect(
        url_for('admin')
    )


@app.route(
    '/admin/user/<int:uid>/toggle',
    methods=['POST']
)
@admin_required
def toggle_user(uid):

    c = db_conn()

    c.execute(
        '''
        UPDATE users
        SET active=CASE active
            WHEN 1 THEN 0
            ELSE 1
        END
        WHERE id=?
        ''',
        (uid,)
    )

    c.commit()
    c.close()

    return redirect(
        url_for('admin')
    )


@app.route(
    '/admin/mix/<mix_id>',
    methods=['POST']
)
@admin_required
def mix_baseline(mix_id):

    v = request.form.get(
        'baseline_mpa',
        ''
    ).strip()

    val = float(v) if v else None

    c = db_conn()

    c.execute(
        'UPDATE mixes SET baseline_mpa=? WHERE mix_id=?',
        (val, mix_id)
    )

    c.commit()
    c.close()

    flash(
        f'{mix_id} baseline updated.',
        'ok'
    )

    return redirect(
        url_for('admin')
    )


if __name__ == '__main__':
    app.run(
        host='0.0.0.0',
        port=int(os.environ.get('PORT', 5000)),
        debug=False
    )
